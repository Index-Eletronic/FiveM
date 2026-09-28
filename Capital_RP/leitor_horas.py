import csv
import os
import re
import shutil
import time
from datetime import datetime, timedelta

try:
    import keyboard
    import pyautogui
    import pytesseract
    from PIL import ImageOps, ImageEnhance
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError as exc:
    print("\nFalta uma dependencia:", exc)
    print("\nInstale com:")
    print("pip install pyautogui keyboard pillow pytesseract openpyxl")
    input("\nPressione ENTER para sair...")
    raise SystemExit(1)


# ============================================================
# CONFIGURACAO GERAL
# ============================================================

META_PADRAO = 42

# Scroll negativo = descer a lista.
# Pequeno de proposito: queremos SOBREPOSICAO entre telas para nao pular membros.
SCROLL_PASSOS = -4

# Se chegarmos ao fim da lista, a tela para de mudar.
# Depois de N varreduras sem nenhum ID novo, encerramos.
MAX_VARREDURAS_SEM_NOVOS = 4

TEMPO_APOS_CLICAR_MEMBRO = 0.45
TEMPO_APOS_ABRIR_DETALHES = 0.70
TEMPO_APOS_FECHAR = 0.30
TEMPO_APOS_SCROLL = 0.90

ARQUIVO_CSV = "membros_pmc.csv"
ARQUIVO_XLSX = "membros_pmc.xlsx"
PASTA_ERROS = "pmc_revisar"

pyautogui.FAILSAFE = True
pyautogui.PAUSE = 0.12

parada_solicitada = False


# ============================================================
# UTILIDADES
# ============================================================

def solicitar_parada():
    global parada_solicitada
    parada_solicitada = True
    print("\n>>> PARADA SEGURA SOLICITADA. Finalizando a operacao atual...")


def localizar_tesseract():
    candidatos = [
        shutil.which("tesseract"),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]

    for caminho in candidatos:
        if caminho and os.path.isfile(caminho):
            return caminho

    return None


def esperar_posicao(atalho, mensagem):
    print("\n" + mensagem)
    print(f"Posicione o mouse e pressione {atalho.upper()}.")

    keyboard.wait(atalho)
    time.sleep(0.35)

    p = pyautogui.position()
    print(f"Registrado: X={p.x} | Y={p.y}")
    return p.x, p.y


def criar_regiao(p1, p2):
    x1, y1 = p1
    x2, y2 = p2

    esquerda = min(x1, x2)
    topo = min(y1, y2)
    direita = max(x1, x2)
    baixo = max(y1, y2)

    return (
        esquerda,
        topo,
        direita - esquerda,
        baixo - topo,
    )


def preparar_imagem(img, escala=3):
    img = ImageOps.grayscale(img)
    img = ImageOps.autocontrast(img)
    img = ImageOps.invert(img)
    img = img.resize((img.width * escala, img.height * escala))
    img = ImageEnhance.Contrast(img).enhance(1.7)
    img = ImageEnhance.Sharpness(img).enhance(1.5)
    return img


def ocr_regiao(regiao, psm=6, salvar_original=None):
    img = pyautogui.screenshot(region=regiao)

    if salvar_original:
        img.save(salvar_original)

    img = preparar_imagem(img)

    return pytesseract.image_to_string(
        img,
        config=f"--psm {psm}",
    )


# ============================================================
# OCR / INTERPRETACAO
# ============================================================

def limpar_nome(nome):
    nome = nome.replace("|", " ")
    nome = nome.strip(" :-_\t\r\n")
    nome = re.sub(r"\s{2,}", " ", nome)
    return nome


def extrair_identidade(texto):
    """
    Espera uma area parecida com:

        Passaporte: 8526
        Name: Kauan Santos
    """
    texto = texto.replace("“", "").replace("”", "")

    passaporte = None
    nome = None

    m = re.search(
        r"Passa?porte\s*[:\-]?\s*(\d{3,8})",
        texto,
        flags=re.IGNORECASE,
    )
    if m:
        passaporte = m.group(1)

    # Fallback: primeiro numero isolado de 3 a 8 digitos.
    if not passaporte:
        numeros = re.findall(r"\b\d{3,8}\b", texto)
        if numeros:
            passaporte = numeros[0]

    m = re.search(
        r"(?:Name|Nome)\s*[:\-]?\s*([^\r\n]+)",
        texto,
        flags=re.IGNORECASE,
    )
    if m:
        nome = limpar_nome(m.group(1))

    # Fallback por linhas.
    if not nome:
        linhas = [
            limpar_nome(linha)
            for linha in texto.splitlines()
            if limpar_nome(linha)
        ]

        for i, linha in enumerate(linhas):
            if re.search(r"\b(?:Name|Nome)\b", linha, flags=re.IGNORECASE):
                resto = re.sub(
                    r".*?(?:Name|Nome)\s*[:\-]?",
                    "",
                    linha,
                    flags=re.IGNORECASE,
                ).strip()

                if resto:
                    nome = limpar_nome(resto)
                    break

                if i + 1 < len(linhas):
                    nome = limpar_nome(linhas[i + 1])
                    break

    if nome and len(nome) < 2:
        nome = None

    return passaporte, nome


def normalizar_hora(valor):
    return valor.replace(".", ":")


def extrair_detalhes(texto):
    """
    Espera algo como:

    Ultimo ponto: 12/09/2026 - 17:03:19
    Total de horas (Semana): 02:43:28
    Total de horas registradas: 86:49:04
    """
    padrao_data = re.compile(
        r"(\d{2}[/-]\d{2}[/-]\d{4}).{0,12}?(\d{1,2}[:.]\d{2}[:.]\d{2})",
        re.DOTALL,
    )

    ultimo_ponto = None
    hora_ultimo_ponto = None

    m = padrao_data.search(texto)

    if m:
        data = m.group(1).replace("-", "/")
        hora_ultimo_ponto = normalizar_hora(m.group(2))
        ultimo_ponto = f"{data} - {hora_ultimo_ponto}"

    horas = re.findall(
        r"\b\d{1,3}[:.]\d{2}[:.]\d{2}\b",
        texto,
    )
    horas = [normalizar_hora(h) for h in horas]

    # Remove a hora que pertence ao "Ultimo ponto".
    if hora_ultimo_ponto:
        removido = False
        restantes = []

        for h in horas:
            if not removido and h == hora_ultimo_ponto:
                removido = True
                continue
            restantes.append(h)

        horas = restantes

    horas_semana = horas[0] if len(horas) >= 1 else None
    horas_total = horas[1] if len(horas) >= 2 else None

    return ultimo_ponto, horas_semana, horas_total


# ============================================================
# HORAS / RELATORIO
# ============================================================

def hora_para_segundos(valor):
    if not valor:
        return 0

    try:
        h, m, s = map(int, valor.split(":"))
        return h * 3600 + m * 60 + s
    except Exception:
        return 0


def segundos_para_excel(segundos):
    # Excel armazena tempo como fracao de um dia.
    return segundos / 86400.0


def salvar_csv(registros):
    campos = [
        "passaporte",
        "nome",
        "ultimo_ponto",
        "horas_semana",
        "horas_total",
        "status",
    ]

    with open(
        ARQUIVO_CSV,
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as f:
        writer = csv.DictWriter(
            f,
            fieldnames=campos,
            delimiter=";",
        )
        writer.writeheader()
        writer.writerows(registros)


def salvar_xlsx(registros):
    wb = Workbook()

    ws = wb.active
    ws.title = "Membros"

    cabecalhos = [
        "PASSAPORTE",
        "NOME",
        "ULTIMO PONTO",
        "HORAS SEMANA",
        "HORAS TOTAL",
        "STATUS",
    ]

    ws.append(cabecalhos)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in registros:
        ws.append([
            r["passaporte"],
            r["nome"],
            r["ultimo_ponto"],
            segundos_para_excel(hora_para_segundos(r["horas_semana"])),
            segundos_para_excel(hora_para_segundos(r["horas_total"])),
            r["status"],
        ])

    # Formato que suporta mais de 24 horas.
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=4).number_format = "[h]:mm:ss"
        ws.cell(row=row, column=5).number_format = "[h]:mm:ss"

    larguras = {
        "A": 14,
        "B": 28,
        "C": 24,
        "D": 18,
        "E": 18,
        "F": 12,
    }

    for col, largura in larguras.items():
        ws.column_dimensions[col].width = largura

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Resumo
    resumo = wb.create_sheet("Resumo")

    total_semana = sum(
        hora_para_segundos(r["horas_semana"])
        for r in registros
    )

    total_geral = sum(
        hora_para_segundos(r["horas_total"])
        for r in registros
    )

    validos_semana = [
        r for r in registros
        if hora_para_segundos(r["horas_semana"]) > 0
    ]

    resumo["A1"] = "RESUMO PMC"
    resumo["A1"].font = Font(bold=True)

    resumo["A3"] = "Membros coletados"
    resumo["B3"] = len(registros)

    resumo["A4"] = "Total de horas - Semana"
    resumo["B4"] = segundos_para_excel(total_semana)
    resumo["B4"].number_format = "[h]:mm:ss"

    resumo["A5"] = "Total de horas registradas"
    resumo["B5"] = segundos_para_excel(total_geral)
    resumo["B5"].number_format = "[h]:mm:ss"

    resumo["A6"] = "Media semanal por membro"
    if registros:
        resumo["B6"] = segundos_para_excel(total_semana / len(registros))
        resumo["B6"].number_format = "[h]:mm:ss"

    resumo["A7"] = "Membros com horas na semana"
    resumo["B7"] = len(validos_semana)

    resumo.column_dimensions["A"].width = 32
    resumo.column_dimensions["B"].width = 18

    # Ranking semanal
    ranking = wb.create_sheet("Ranking Semana")
    ranking.append(["POSICAO", "PASSAPORTE", "NOME", "HORAS SEMANA"])

    for cell in ranking[1]:
        cell.font = Font(bold=True)

    ordenados = sorted(
        registros,
        key=lambda x: hora_para_segundos(x["horas_semana"]),
        reverse=True,
    )

    for posicao, r in enumerate(ordenados, start=1):
        ranking.append([
            posicao,
            r["passaporte"],
            r["nome"],
            segundos_para_excel(hora_para_segundos(r["horas_semana"])),
        ])

    for row in range(2, ranking.max_row + 1):
        ranking.cell(row=row, column=4).number_format = "[h]:mm:ss"

    ranking.column_dimensions["A"].width = 12
    ranking.column_dimensions["B"].width = 14
    ranking.column_dimensions["C"].width = 28
    ranking.column_dimensions["D"].width = 18

    wb.save(ARQUIVO_XLSX)


def salvar_tudo(registros):
    salvar_csv(registros)
    salvar_xlsx(registros)


# ============================================================
# CALIBRACAO
# ============================================================

def calibrar():
    print("""
================================================================
CALIBRACAO DO COLETOR PMC
================================================================

Deixe a tela MEMBROS aberta e o PRIMEIRO membro visivel.

IMPORTANTE:
Quando eu pedir uma linha, posicione o mouse sobre o NOME do membro.
Nao precisa clicar; apenas use o atalho indicado.
""")

    input("Pressione ENTER e volte para o FiveM...")

    primeira_linha = esperar_posicao(
        "ctrl+alt+1",
        "1/10 - Posicao do NOME do PRIMEIRO membro visivel"
    )

    segunda_linha = esperar_posicao(
        "ctrl+alt+2",
        "2/10 - Posicao do NOME do SEGUNDO membro visivel"
    )

    ultima_linha = esperar_posicao(
        "ctrl+alt+3",
        "3/10 - Posicao do NOME do ULTIMO membro totalmente visivel"
    )

    # Seleciona primeiro membro para exibir Passaporte/Name no painel esquerdo.
    print("\nSelecionando o primeiro membro...")
    pyautogui.click(*primeira_linha)
    time.sleep(TEMPO_APOS_CLICAR_MEMBRO)

    botao_detalhes = esperar_posicao(
        "ctrl+alt+4",
        "4/10 - Centro do botao DETALHES no painel esquerdo"
    )

    identidade_tl = esperar_posicao(
        "ctrl+alt+5",
        "5/10 - Canto SUPERIOR ESQUERDO da area Passaporte/Name"
    )

    identidade_br = esperar_posicao(
        "ctrl+alt+6",
        "6/10 - Canto INFERIOR DIREITO da area Passaporte/Name"
    )

    regiao_identidade = criar_regiao(
        identidade_tl,
        identidade_br,
    )

    print("\nAbrindo Detalhes para calibrar a janela...")
    pyautogui.click(*botao_detalhes)
    time.sleep(TEMPO_APOS_ABRIR_DETALHES)

    detalhes_tl = esperar_posicao(
        "ctrl+alt+7",
        "7/10 - Canto SUPERIOR ESQUERDO da janela Detalhes"
    )

    detalhes_br = esperar_posicao(
        "ctrl+alt+8",
        "8/10 - Canto INFERIOR DIREITO da janela Detalhes"
    )

    botao_fechar = esperar_posicao(
        "ctrl+alt+9",
        "9/10 - Centro do botao FECHAR"
    )

    regiao_detalhes = criar_regiao(
        detalhes_tl,
        detalhes_br,
    )

    pyautogui.click(*botao_fechar)
    time.sleep(TEMPO_APOS_FECHAR)

    area_scroll = esperar_posicao(
        "ctrl+alt+0",
        "10/10 - Um ponto vazio no MEIO DA LISTA para realizar a rolagem"
    )

    # Calcula linhas visiveis
    espacamento = segunda_linha[1] - primeira_linha[1]

    if abs(espacamento) < 20:
        raise RuntimeError(
            "O espacamento entre a primeira e segunda linha ficou muito pequeno."
        )

    passo = abs(espacamento)
    sentido = 1 if ultima_linha[1] >= primeira_linha[1] else -1

    linhas = []
    y = primeira_linha[1]

    while True:
        linhas.append((primeira_linha[0], int(y)))

        if sentido == 1:
            if y + passo > ultima_linha[1] + passo * 0.30:
                break
            y += passo
        else:
            if y - passo < ultima_linha[1] - passo * 0.30:
                break
            y -= passo

        if len(linhas) > 30:
            break

    print("\nCALIBRACAO CONCLUIDA")
    print(f"Espacamento vertical: {passo}px")
    print(f"Linhas visiveis: {len(linhas)}")

    return {
        "linhas": linhas,
        "botao_detalhes": botao_detalhes,
        "botao_fechar": botao_fechar,
        "regiao_identidade": regiao_identidade,
        "regiao_detalhes": regiao_detalhes,
        "area_scroll": area_scroll,
    }


# ============================================================
# COLETA
# ============================================================

def processar_slot(config, indice, registros, ids_processados):
    """
    Retorna True quando encontrou um NOVO membro.
    """
    x, y = config["linhas"][indice]

    # Seleciona membro.
    pyautogui.click(x, y)
    time.sleep(TEMPO_APOS_CLICAR_MEMBRO)

    # Le identidade no painel esquerdo.
    texto_id = ocr_regiao(
        config["regiao_identidade"],
        psm=6,
    )

    passaporte, nome = extrair_identidade(texto_id)

    if not passaporte:
        print(
            f"Slot {indice + 1}: sem ID reconhecido "
            "(linha vazia ou OCR falhou)."
        )
        return False

    if passaporte in ids_processados:
        print(
            f"Slot {indice + 1}: {passaporte} - "
            f"{nome or '?'} [JA COLETADO]"
        )
        return False

    print(
        f"\n>>> NOVO: {passaporte} - "
        f"{nome or 'NOME NAO LIDO'}"
    )

    # Abre Detalhes.
    pyautogui.click(*config["botao_detalhes"])
    time.sleep(TEMPO_APOS_ABRIR_DETALHES)

    texto_detalhes = ocr_regiao(
        config["regiao_detalhes"],
        psm=6,
    )

    ultimo_ponto, horas_semana, horas_total = extrair_detalhes(
        texto_detalhes
    )

    status = "OK"

    if not all([
        passaporte,
        nome,
        ultimo_ponto,
        horas_semana,
        horas_total,
    ]):
        status = "REVISAR"

        os.makedirs(PASTA_ERROS, exist_ok=True)

        # Guarda somente os casos problemáticos.
        imagem_erro = os.path.join(
            PASTA_ERROS,
            f"{passaporte}_detalhes.png",
        )

        pyautogui.screenshot(
            region=config["regiao_detalhes"]
        ).save(imagem_erro)

    registro = {
        "passaporte": passaporte,
        "nome": nome or "",
        "ultimo_ponto": ultimo_ponto or "",
        "horas_semana": horas_semana or "",
        "horas_total": horas_total or "",
        "status": status,
    }

    registros.append(registro)
    ids_processados.add(passaporte)

    # CHECKPOINT a cada membro.
    salvar_tudo(registros)

    print(f"Ultimo ponto : {registro['ultimo_ponto'] or 'NAO LIDO'}")
    print(f"Horas semana : {registro['horas_semana'] or 'NAO LIDO'}")
    print(f"Horas total  : {registro['horas_total'] or 'NAO LIDO'}")
    print(f"Status       : {status}")
    print(f"Coletados    : {len(registros)}")

    # Fecha popup.
    pyautogui.click(*config["botao_fechar"])
    time.sleep(TEMPO_APOS_FECHAR)

    return True


def main():
    global parada_solicitada

    print("=" * 68)
    print("COLETOR AUTOMATICO PMC - V1.0")
    print("ID + Nome + Ultimo Ponto + Horas Semana + Horas Total")
    print("=" * 68)

    tesseract = localizar_tesseract()

    if not tesseract:
        print("\nTesseract OCR nao encontrado.")
        print(r"Esperado em C:\Program Files\Tesseract-OCR\tesseract.exe")
        input("\nPressione ENTER para sair...")
        return

    pytesseract.pytesseract.tesseract_cmd = tesseract

    print(f"\nTesseract encontrado: {tesseract}")

    print("""
SEGURANCA
---------
CTRL + ALT + Q
    Solicita parada segura.

Mover o mouse rapidamente para o CANTO SUPERIOR ESQUERDO da tela
    Aciona o FAILSAFE do PyAutoGUI e interrompe o robo.
""")

    keyboard.add_hotkey(
        "ctrl+alt+q",
        solicitar_parada,
    )

    registros = []

    try:
        entrada = input(
            f"Quantidade de membros para coletar [{META_PADRAO}]: "
        ).strip()

        meta = int(entrada) if entrada else META_PADRAO

        config = calibrar()

        print("""
================================================================
PRONTO PARA COLETAR
================================================================

O robo vai:

1. Clicar em cada membro visivel.
2. Ler ID/Passaporte + Nome.
3. Abrir Detalhes.
4. Ler Ultimo ponto + Horas da Semana + Horas Totais.
5. Fechar Detalhes.
6. Ir para o proximo.
7. Rolar a lista.
8. Ignorar IDs repetidos.
9. Parar ao atingir a meta ou chegar ao fim da lista.

Arquivos gerados:
    membros_pmc.csv
    membros_pmc.xlsx

Comecando em 5 segundos.
NAO mexa no mouse durante a coleta.
""")

        for i in range(5, 0, -1):
            print(i)
            time.sleep(1)

        ids_processados = set()
        varreduras_sem_novos = 0
        numero_varredura = 0

        salvar_tudo(registros)

        while (
            not parada_solicitada
            and len(registros) < meta
            and varreduras_sem_novos < MAX_VARREDURAS_SEM_NOVOS
        ):
            numero_varredura += 1
            quantidade_inicio = len(registros)

            print("\n" + "=" * 68)
            print(
                f"VARREDURA {numero_varredura} | "
                f"{len(registros)}/{meta} membros"
            )
            print("=" * 68)

            for indice in range(len(config["linhas"])):
                if parada_solicitada:
                    break

                if len(registros) >= meta:
                    break

                processar_slot(
                    config,
                    indice,
                    registros,
                    ids_processados,
                )

            novos = len(registros) - quantidade_inicio

            if novos == 0:
                varreduras_sem_novos += 1
                print(
                    f"\nNenhum ID novo nesta tela "
                    f"({varreduras_sem_novos}/"
                    f"{MAX_VARREDURAS_SEM_NOVOS})."
                )
            else:
                varreduras_sem_novos = 0
                print(f"\nNovos membros nesta tela: {novos}")

            if (
                parada_solicitada
                or len(registros) >= meta
                or varreduras_sem_novos >= MAX_VARREDURAS_SEM_NOVOS
            ):
                break

            # Scroll com sobreposicao.
            print("\nDescendo a lista...")
            pyautogui.moveTo(
                *config["area_scroll"],
                duration=0.15,
            )

            pyautogui.scroll(SCROLL_PASSOS)
            time.sleep(TEMPO_APOS_SCROLL)

        salvar_tudo(registros)

        print("\n" + "=" * 68)
        print("COLETA ENCERRADA")
        print("=" * 68)
        print(f"Membros coletados : {len(registros)}")
        print(f"Meta               : {meta}")
        print(f"CSV                : {os.path.abspath(ARQUIVO_CSV)}")
        print(f"Excel              : {os.path.abspath(ARQUIVO_XLSX)}")

        revisar = [
            r for r in registros
            if r["status"] != "OK"
        ]

        print(f"Registros REVISAR  : {len(revisar)}")

        if revisar:
            print("\nIDs para revisar:")
            print(
                ", ".join(
                    r["passaporte"]
                    for r in revisar
                )
            )
            print(
                f"Capturas salvas em: "
                f"{os.path.abspath(PASTA_ERROS)}"
            )

        if len(registros) >= meta:
            print("\n>>> META ATINGIDA COM SUCESSO.")

        elif parada_solicitada:
            print("\n>>> Coleta interrompida pelo usuario.")

        else:
            print(
                "\n>>> O robo nao encontrou IDs novos por varias "
                "varreduras e assumiu que chegou ao fim."
            )

    except pyautogui.FailSafeException:
        print(
            "\nFAILSAFE acionado. "
            "Automacao interrompida imediatamente."
        )

    except KeyboardInterrupt:
        print("\nCancelado pelo teclado.")

    except Exception as exc:
        print(f"\nERRO: {exc}")

    finally:
        try:
            salvar_tudo(registros)
        except Exception:
            pass

        keyboard.unhook_all_hotkeys()

    input("\nPressione ENTER para encerrar...")


if __name__ == "__main__":
    main()
